"""Write the scripted-model fixtures for three real Toolathlon-Gym tasks. Run from the repo root.

Each task has one plan: read-only calls (the researcher's), write calls (the
operator's), and the claims each sub-agent returns. The plan is rendered three
ways: ``sa`` (the main agent makes every call), ``vanilla_mad`` (sub-agents
reply in free text) and ``evidence_risk`` (sub-agents reply with claims that
cite their call ids). Spans and values were checked against the gym ed735ba
database; the scripts solve the tasks, so the gym's own evaluation passes.

Method counts are ``derived`` claims: a count of list items does not appear
verbatim in the tool output, so L1 does not look for it.

``sa_lazy.json`` claims done without doing anything; its evaluation must fail.
``evidence_reinject.json`` (sf-hr-rating-distribution only) gives the researcher
a second round with fresh call ids, for a run where the fault injector corrupts
the first return and the audit sends it back.
"""

from __future__ import annotations

import json
from pathlib import Path

HERE = Path(__file__).resolve().parent
WS = "/workspace/agent_workspace"
TICKERS = ["GOOGL", "AMZN", "JPM", "JNJ", "XOM"]

# Counts from the HR employees table (sf_data.HR_ANALYTICS__PUBLIC__EMPLOYEES).
HR_COUNTS = {
    "Engineering": [375, 1006, 3180, 1814, 721],
    "Finance": [334, 1053, 3251, 1789, 721],
    "HR": [358, 1067, 3167, 1745, 740],
    "Operations": [380, 1051, 3263, 1739, 687],
    "R&D": [362, 1053, 3190, 1785, 693],
    "Sales": [373, 1121, 3215, 1829, 694],
    "Support": [371, 1108, 3250, 1763, 752],
}
HR_QUERY = (
    'SELECT "DEPARTMENT", "PERFORMANCE_RATING", COUNT(*) AS n '
    'FROM sf_data."HR_ANALYTICS__PUBLIC__EMPLOYEES" GROUP BY 1, 2 ORDER BY 1, 2'
)

PORTFOLIO_CODE = """\
import math, os
from collections import defaultdict
import openpyxl, psycopg2

tickers = ["GOOGL", "AMZN", "JPM", "JNJ", "XOM"]
conn = psycopg2.connect(host=os.environ["PGHOST"], dbname="toolathlon_gym", user="eigent", password="camel")
cur = conn.cursor()
prices = defaultdict(dict)
for t in tickers:
    cur.execute("SELECT date, close FROM yf.stock_prices WHERE symbol = %s ORDER BY date", (t,))
    for day, close in cur.fetchall():
        prices[day][t] = float(close)
dates = sorted(d for d, row in prices.items() if len(row) == len(tickers))
wb = openpyxl.Workbook()
history = wb.active
history.title = "Price History"
history.append(["Date", *tickers])
for d in dates:
    history.append([d.isoformat(), *[round(prices[d][t], 2) for t in tickers]])
metrics = wb.create_sheet("Risk Metrics")
metrics.append(["Ticker", "Average Close Price", "Standard Deviation", "Minimum Close", "Maximum Close"])
assessment = wb.create_sheet("Risk Assessment")
assessment.append(["Ticker", "Coefficient of Variation", "Risk Category"])
for t in tickers:
    values = [prices[d][t] for d in dates]
    mean = sum(values) / len(values)
    std = math.sqrt(sum((v - mean) ** 2 for v in values) / (len(values) - 1))
    cv = std / mean
    category = "Low Risk" if cv < 0.10 else ("Medium Risk" if cv <= 0.20 else "High Risk")
    metrics.append([t, round(mean, 2), round(std, 2), round(min(values), 2), round(max(values), 2)])
    assessment.append([t, round(cv, 4), category])
wb.save("/workspace/agent_workspace/Portfolio_Risk_Analysis.xlsx")
print(f"saved Portfolio_Risk_Analysis.xlsx with {len(dates)} dates")
"""


def call(tool: str, arguments: dict, call_id: str) -> dict:
    return {"tool": tool, "arguments": arguments, "call_id": call_id}


def claim(claim_id: str, text: str, kind: str, value: object, evidence: list[tuple[str, str]], depends_on: list[str] | None = None) -> dict:
    return {
        "claim_id": claim_id,
        "confidence": "high",
        "depends_on": depends_on or [],
        "evidence": [{"call_id": cid, "span": span, "type": "tool_call"} for cid, span in evidence],
        "kind": kind,
        "text": text,
        "value": value,
    }


def wc_plan() -> dict:
    path = f"{WS}/WC_Shipping_Report.xlsx"
    zones = [("California", 2, 2), ("Domestic US", 1, 3), ("International", 3, 1)]
    reads = [call("woocommerce-woo_shipping_zones_list", {}, "r_zones")]
    reads += [call("woocommerce-woo_shipping_zone_methods_list", {"zoneId": zid}, f"r_methods_{zid}") for _, zid, _ in zones]
    writes = [
        call("excel-create_workbook", {"filepath": path}, "w_book"),
        call(
            "excel-write_data_to_excel",
            {"filepath": path, "sheet_name": "Shipping Zones", "data": [["Zone", "Shipping_Methods"], *[[n, c] for n, _, c in zones]]},
            "w_zones",
        ),
        call(
            "excel-write_data_to_excel",
            {"filepath": path, "sheet_name": "Summary", "data": [["Metric", "Value"], ["Total_Zones", 3], ["Total_Methods", 6]]},
            "w_summary",
        ),
        call(
            "emails-send_email",
            {
                "to": "logistics@shop.com",
                "subject": "Shipping Zone Configuration Review",
                "body": "Shipping zones and method counts:\n- California: 2\n- Domestic US: 3\n- International: 1\nTotal: 3 zones, 6 methods.",
            },
            "w_email",
        ),
    ]
    facts = [
        claim("c1", "The shop has three shipping zones: International, California and Domestic US.", "fact",
              ["International", "California", "Domestic US"], [("r_zones", '"name": "International"')]),
        claim("c2", "Domestic US (zone 1) has 3 shipping methods.", "derived", 3, [("r_methods_1", '"method_title": "Standard Shipping"')]),
        claim("c3", "California (zone 2) has 2 shipping methods.", "derived", 2, [("r_methods_2", '"method_title": "CA Standard"')]),
        claim("c4", "International (zone 3) has 1 shipping method.", "derived", 1, [("r_methods_3", '"method_title": "International Shipping"')]),
        claim("c5", "The three zones have 6 shipping methods in total.", "derived", 6,
              [("r_methods_1", '"zone_id": 1'), ("r_methods_2", '"zone_id": 2'), ("r_methods_3", '"zone_id": 3')], ["c2", "c3", "c4"]),
    ]
    acts = [
        claim("c1", "Wrote the Shipping Zones sheet of WC_Shipping_Report.xlsx.", "artifact", "Domestic US",
              [("w_zones", "Data written to Shipping Zones")]),
        claim("c2", "Wrote the Summary sheet with Total_Zones 3 and Total_Methods 6.", "artifact", "Total_Methods",
              [("w_summary", "Data written to Summary")]),
        claim("c3", "Emailed the zone list to logistics@shop.com.", "action_done", "Shipping Zone Configuration Review",
              [("w_email", "Email sent successfully to logistics@shop.com")]),
    ]
    return {
        "reads": reads, "writes": writes, "facts": facts, "acts": acts,
        "research_input": "List every shipping zone in the shop and how many shipping methods each zone has.",
        "operate_input": (
            f"Create {path} with a sheet 'Shipping Zones' (columns Zone, Shipping_Methods; rows California 2, "
            "Domestic US 3, International 1, sorted by zone) and a sheet 'Summary' (Metric, Value; Total_Zones 3, "
            "Total_Methods 6). Then email logistics@shop.com with subject 'Shipping Zone Configuration Review' listing the zones."
        ),
        "research_text": "Zones: California (2 methods), Domestic US (3 methods), International (1 method). 6 methods in total.",
        "operate_text": "Wrote both sheets of WC_Shipping_Report.xlsx and sent the review email to logistics@shop.com.",
    }


def hr_plan() -> dict:
    path = f"{WS}/HR_Rating_Distribution.xlsx"
    rows = [[f"{dept} - {rating}", counts[rating - 1]] for dept, counts in sorted(HR_COUNTS.items()) for rating in range(1, 6)]
    by_rating = [sum(counts[i] for counts in HR_COUNTS.values()) for i in range(5)]
    total = sum(by_rating)
    high = round(100 * (by_rating[3] + by_rating[4]) / total, 1)
    summary = [["Metric", "Value"], ["Total_Employees", total], ["Rating_5_Count", by_rating[4]],
               ["Rating_4_Count", by_rating[3]], ["Rating_1_Count", by_rating[0]], ["High_Performers_Pct", high]]
    reads = [call("snowflake-read_query", {"query": HR_QUERY}, "r_counts")]
    writes = [
        call("excel-create_workbook", {"filepath": path}, "w_book"),
        call("excel-write_data_to_excel", {"filepath": path, "sheet_name": "Rating Distribution", "data": [["Dept_Rating", "Count"], *rows]}, "w_dist"),
        call("excel-write_data_to_excel", {"filepath": path, "sheet_name": "Summary", "data": summary}, "w_summary"),
        call(
            "emails-send_email",
            {
                "to": "hr-analytics@company.com",
                "subject": "Performance Rating Distribution Report",
                "body": (
                    f"Performance rating distribution across {total} employees: {by_rating[4]} rated 5, "
                    f"{by_rating[3]} rated 4, {by_rating[0]} rated 1. High performers (rating 4 or 5): {high}%."
                ),
            },
            "w_email",
        ),
    ]
    facts = [
        claim("c1", "Engineering has 721 employees with performance rating 5.", "fact", 721, [("r_counts", "N: 721")]),
        claim("c2", "Support has 752 employees with performance rating 5.", "fact", 752, [("r_counts", "N: 752")]),
        claim("c3", f"Across all departments {total} employees are rated; {high}% have rating 4 or 5.", "derived", high,
              [("r_counts", "DEPARTMENT: Engineering"), ("r_counts", "DEPARTMENT: Support")], ["c1", "c2"]),
    ]
    acts = [
        claim("c1", "Wrote the Rating Distribution sheet of HR_Rating_Distribution.xlsx.", "artifact", "Engineering - 5",
              [("w_dist", "Data written to Rating Distribution")]),
        claim("c2", "Wrote the Summary sheet.", "artifact", "High_Performers_Pct", [("w_summary", "Data written to Summary")]),
        claim("c3", "Emailed the distribution summary to hr-analytics@company.com.", "action_done",
              "Performance Rating Distribution Report", [("w_email", "Email sent successfully to hr-analytics@company.com")]),
    ]
    return {
        "reads": reads, "writes": writes, "facts": facts, "acts": acts,
        "research_input": "Count employees by department and performance rating in the HR data warehouse.",
        "operate_input": (
            f"Create {path} with the 'Rating Distribution' and 'Summary' sheets from these counts, then email "
            "hr-analytics@company.com with subject 'Performance Rating Distribution Report'."
        ),
        "research_text": f"Counts per department and rating retrieved; {total} employees, {high}% rated 4 or 5.",
        "operate_text": "Wrote HR_Rating_Distribution.xlsx and sent the summary email.",
    }


def yf_plan() -> dict:
    first_close = {"GOOGL": "130.4146", "AMZN": "173.51", "JPM": "181.5711", "JNJ": "150.2022", "XOM": "99.8529"}
    reads = [call("yahoo-finance-get_historical_stock_prices", {"ticker": t, "period": "max"}, f"r_{t.lower()}") for t in TICKERS]
    writes = [call("python_execute", {"code": PORTFOLIO_CODE, "filename": "portfolio_risk.py"}, "w_python")]
    facts = [
        claim(f"c{i}", f"{t} has closing prices from 2024-03-06 to 2026-03-05; the first close is {first_close[t]}.", "fact",
              float(first_close[t]), [(f"r_{t.lower()}", '"Date":"2024-03-06T00:00:00.000"')])
        for i, t in enumerate(TICKERS, start=1)
    ]
    acts = [
        claim("c1", "Saved Portfolio_Risk_Analysis.xlsx with Price History, Risk Metrics and Risk Assessment sheets.", "artifact",
              "Portfolio_Risk_Analysis.xlsx", [("w_python", "saved Portfolio_Risk_Analysis.xlsx with 501 dates")]),
    ]
    return {
        "reads": reads, "writes": writes, "facts": facts, "acts": acts,
        "research_input": "Retrieve the full daily closing price history for GOOGL, AMZN, JPM, JNJ and XOM.",
        "operate_input": (
            f"Build {WS}/Portfolio_Risk_Analysis.xlsx: 'Price History' (Date plus one close column per ticker, dates where "
            "all five have data), 'Risk Metrics' (mean, sample standard deviation, min, max, two decimals) and "
            "'Risk Assessment' (coefficient of variation to four decimals and Low/Medium/High Risk at 0.10 and 0.20)."
        ),
        "research_text": "All five tickers have 501 daily closes from 2024-03-06 to 2026-03-05.",
        "operate_text": "Saved Portfolio_Risk_Analysis.xlsx with the three sheets.",
    }


PLANS = {
    "wc-shipping-analysis": wc_plan,
    "sf-hr-rating-distribution": hr_plan,
    "yf-portfolio-risk-excel": yf_plan,
}

DONE = call("claim_done", {}, "m_done")


def subagent_return(subagent_id: str, claims: list[dict]) -> str:
    return json.dumps(
        {"claims": claims, "open_issues": [], "status": "complete", "subagent_id": subagent_id}, sort_keys=True
    )


def render(plan: dict, arm: str) -> dict:
    if arm == "sa":
        return {"agents": {"main": [*plan["reads"], *plan["writes"], DONE]}}
    structured = arm == "evidence_risk"
    research_reply = subagent_return("researcher", plan["facts"]) if structured else plan["research_text"]
    operate_reply = subagent_return("operator", plan["acts"]) if structured else plan["operate_text"]
    return {
        "agents": {
            "main": [
                call("researcher", {"input": plan["research_input"]}, "m_research"),
                call("operator", {"input": plan["operate_input"]}, "m_operate"),
                DONE,
            ],
            "researcher": [*plan["reads"], {"content": research_reply}],
            "operator": [*plan["writes"], {"content": operate_reply}],
        }
    }


def main() -> None:
    usage = {"input_tokens": 1200, "cached_tokens": 900, "output_tokens": 150}
    for arm in ("sa", "vanilla_mad", "evidence_risk"):
        payload = {
            "served_model": "mock-scripted",
            "usage": usage,
            "tasks": {name: render(build(), arm) for name, build in PLANS.items()},
        }
        (HERE / f"{arm}.json").write_text(json.dumps(payload, indent=1, sort_keys=True) + "\n", encoding="utf-8")
    hr = hr_plan()
    retry = call("snowflake-read_query", {"query": HR_QUERY}, "r_counts_retry")
    first_facts = [c for c in hr["facts"] if c["kind"] == "fact"]
    retry_facts = [{**c, "evidence": [{**e, "call_id": "r_counts_retry"} for e in c["evidence"]]} for c in first_facts]
    reinject = render(hr, "evidence_risk")
    reinject["agents"]["researcher"] = [
        *hr["reads"], {"content": subagent_return("researcher", first_facts)},
        retry, {"content": subagent_return("researcher", retry_facts)},
    ]
    (HERE / "evidence_reinject.json").write_text(
        json.dumps({"served_model": "mock-scripted", "usage": usage, "tasks": {"sf-hr-rating-distribution": reinject}},
                   indent=1, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    lazy = {"served_model": "mock-scripted", "usage": usage, "tasks": {name: {"agents": {"main": [DONE]}} for name in PLANS}}
    (HERE / "sa_lazy.json").write_text(json.dumps(lazy, indent=1, sort_keys=True) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
